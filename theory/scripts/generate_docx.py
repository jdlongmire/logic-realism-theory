#!/usr/bin/env python3
"""
Generate LRT-MASTER-v2.0.docx from markdown with proper formatting:
- Centered figures and tables
- Tables fit within margins
- Header rows with charcoal background and white text
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Colors as hex strings for shading
CHARCOAL_HEX = '363636'
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CHARCOAL_RGB = RGBColor(0x36, 0x36, 0x36)

def set_cell_shading(cell, hex_color):
    """Set background color of a table cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), hex_color.upper())
    tcPr.append(shd)

def parse_markdown_table(lines):
    """Parse markdown table lines into rows of cells."""
    rows = []
    for line in lines:
        if '|' in line:
            # Skip separator lines (----)
            if re.match(r'^\|[\s\-:|]+\|$', line.strip()):
                continue
            cells = [c.strip() for c in line.strip().strip('|').split('|')]
            if cells and any(c for c in cells):
                rows.append(cells)
    return rows

def add_styled_table(doc, rows, max_width_inches=6.0):
    """Add a table with proper formatting."""
    if not rows:
        return

    num_cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Calculate column width
    col_width = Inches(max_width_inches / num_cols)

    for i, row_data in enumerate(rows):
        row = table.rows[i]
        for j, cell_text in enumerate(row_data):
            if j < num_cols:
                cell = row.cells[j]
                cell.text = cell_text
                cell.width = col_width

                # Style header row (first row)
                if i == 0:
                    set_cell_shading(cell, CHARCOAL_HEX)
                    for paragraph in cell.paragraphs:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for run in paragraph.runs:
                            run.font.color.rgb = WHITE
                            run.font.bold = True
                else:
                    for paragraph in cell.paragraphs:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Add spacing after table
    doc.add_paragraph()
    return table

def add_centered_paragraph(doc, text, style=None):
    """Add a centered paragraph."""
    p = doc.add_paragraph(text, style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p

def process_markdown(md_path, output_path, figures_dir):
    """Convert markdown to styled DOCX."""
    doc = Document()

    # Set default margins
    for section in doc.sections:
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)

    with open(md_path, 'r') as f:
        content = f.read()

    lines = content.split('\n')
    i = 0
    in_table = False
    table_lines = []

    while i < len(lines):
        line = lines[i]

        # Handle tables
        if '|' in line and not line.strip().startswith('!['):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            # End of table
            rows = parse_markdown_table(table_lines)
            if rows:
                add_styled_table(doc, rows)
            in_table = False
            table_lines = []

        # Handle figures (centered)
        if line.strip().startswith('!['):
            match = re.match(r'!\[(.*?)\]\((.*?)\)', line.strip())
            if match:
                caption, img_path = match.groups()
                # Resolve relative path
                if img_path.startswith('../figures/'):
                    img_file = figures_dir / img_path.replace('../figures/', '')
                else:
                    img_file = Path(img_path)

                if img_file.exists() and img_file.suffix.lower() == '.svg':
                    # SVG - add placeholder text (DOCX doesn't support SVG natively)
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run(f'[Figure: {caption}]')
                    run.italic = True
                    # Add caption
                    cap = doc.add_paragraph(caption)
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif img_file.exists():
                    # Other image formats
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run()
                    run.add_picture(str(img_file), width=Inches(5.0))
                    # Add caption
                    cap = doc.add_paragraph(caption)
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    # Image not found
                    p = doc.add_paragraph(f'[Figure: {caption} - not found: {img_path}]')
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # Handle headers
        if line.startswith('# '):
            doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith('#### '):
            doc.add_heading(line[5:].strip(), level=3)
        elif line.strip() == '---':
            # Horizontal rule - add some spacing
            doc.add_paragraph()
        elif line.strip().startswith('> '):
            # Blockquote
            p = doc.add_paragraph(line.strip()[2:])
            p.paragraph_format.left_indent = Inches(0.5)
            p.italic = True
        elif line.strip().startswith('- '):
            # Bullet point
            doc.add_paragraph(line.strip()[2:], style='List Bullet')
        elif line.strip().startswith('$$'):
            # Display math (centered)
            # Collect until closing $$
            math_lines = [line.strip()[2:]]
            i += 1
            while i < len(lines) and '$$' not in lines[i]:
                math_lines.append(lines[i])
                i += 1
            if i < len(lines):
                math_lines.append(lines[i].replace('$$', ''))
            math_text = ' '.join(math_lines).strip()
            p = doc.add_paragraph(math_text)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.italic = True
        elif line.strip():
            # Regular paragraph
            # Handle inline formatting
            text = line.strip()
            # Remove markdown bold/italic for now (would need more complex handling)
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            text = re.sub(r'\*(.+?)\*', r'\1', text)
            doc.add_paragraph(text)

        i += 1

    # Handle any remaining table
    if in_table and table_lines:
        rows = parse_markdown_table(table_lines)
        if rows:
            add_styled_table(doc, rows)

    doc.save(output_path)
    print(f'Generated: {output_path}')

if __name__ == '__main__':
    base_dir = Path('/media/jdlongmire/Macro-Drive-2TB/GitHub_Repos/logic-realism-theory/theory')
    md_path = base_dir / 'LRT-MASTER-v2.0.md'
    output_path = base_dir / 'docs' / 'LRT-MASTER-v2.0.docx'
    figures_dir = base_dir / 'figures'

    # Ensure output directory exists
    output_path.parent.mkdir(parents=True, exist_ok=True)

    process_markdown(md_path, output_path, figures_dir)
